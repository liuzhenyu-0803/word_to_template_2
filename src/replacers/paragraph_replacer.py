"""  
段落替换模块  
负责处理Word文档中段落内容的替换
"""  

import os
import json
import re
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def replace_values_with_placeholders(doc, match_results_dir):
    """
    将Word文档中的段落内容替换为占位符，用于生成模板
    
    参数:
        doc: docx.Document对象，要处理的原始文档
        match_results_dir: 匹配结果目录路径
    """
    value_to_key_map = {}
    # 自动查找所有paragraph_*.json
    match_files = [f for f in os.listdir(match_results_dir) if f.startswith('paragraph_') and f.endswith('_matches.json')]
    if not match_files:
        print("未找到段落匹配结果文件")
        return
    # 加载所有段落匹配结果
    for match_file in match_files:
        file_path = os.path.join(match_results_dir, match_file)
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                match_data = json.load(f)
                for item in match_data:
                    if 'field_key' in item and 'field_value' in item:
                        key = item['field_key']
                        value = item['field_value']
                        if value and value.strip():
                            value_to_key_map[value] = key
        except Exception as e:
            print(f"读取匹配结果文件 {match_file} 失败: {e}")
    print(f"加载了 {len(value_to_key_map)} 个段落值-键映射项")
    sorted_values = sorted(value_to_key_map.keys(), key=len, reverse=True)
    for para in doc.paragraphs:
        para_text = para.text.strip()
        if not para_text:
            continue
        replaced = False
        new_text = para_text
        if para_text in value_to_key_map:
            key = value_to_key_map[para_text]
            new_text = f"{{{{{key}}}}}"
            replaced = True
        else:
            for value in sorted_values:
                if value in para_text:
                    key = value_to_key_map[value]
                    pattern = re.escape(value)
                    new_text = re.sub(pattern, f"{{{{{key}}}}}", new_text)
                    replaced = True
        if replaced and new_text != para_text:
            para.clear()
            para.add_run(new_text)
