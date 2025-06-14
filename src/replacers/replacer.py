"""  
Word文档替换器 - 封装文档元素替换功能  
调用专门的替换器模块完成实际工作  
"""  

import os
from docx import Document
from . import paragraph_replacer
from . import table_replacer

def replace_document(original_doc_path, match_results_dir, template_doc_path):
    """
    根据匹配结果将Word文档中的实际内容替换为占位符，生成模板
    
    参数:
        original_doc_path: 原始Word文档路径
        match_results_dir: 匹配结果目录路径，包含从matcher得到的匹配结果
        template_doc_path: 生成的模板文档输出路径
    """
    # 检查输入文件和目录是否存在
    if not os.path.exists(original_doc_path):
        print(f"错误：原始文档 {original_doc_path} 不存在")
        return
        
    if not os.path.exists(match_results_dir):
        print(f"错误：匹配结果目录 {match_results_dir} 不存在")
        return
    
    # 确保输出目录存在
    output_dir = os.path.dirname(template_doc_path)
    if not os.path.exists(output_dir):
        os.makedirs(output_dir, exist_ok=True)
    
    # 读取原始文档
    print(f"正在读取原始文档: {original_doc_path}")
    doc = Document(original_doc_path)
    
    # 表格内容替换
    table_replacer.replace_values_with_placeholders(doc, match_results_dir)
    
    # 段落内容替换
    # paragraph_replacer.replace_values_with_placeholders(doc, match_results_dir)
    
    # 保存处理后的模板文档
    doc.save(template_doc_path)
    print(f"已保存生成的模板文档: {template_doc_path}")

# 直接运行时的入口点
if __name__ == "__main__":
    project_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    original_doc_path = os.path.join(project_dir, "document/document.docx")
    match_results_dir = os.path.join(project_dir, "document/match_results")
    template_doc_path = os.path.join(project_dir, "document/template.docx")
    
    replace_document(original_doc_path, match_results_dir, template_doc_path)
